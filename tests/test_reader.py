import sys, unittest, threading, json
from pathlib import Path
from io import BytesIO
from urllib.request import Request, urlopen
from urllib.error import HTTPError
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import Relatorio_saidas as app
from openpyxl import load_workbook

def xml(key='1'*44,ret='0.00',namespace=True):
    return f'''<nfeProc {'xmlns="http://www.portalfiscal.inf.br/nfe"' if namespace else ''}><NFe><infNFe Id="NFe{key}"><ide><nNF>000123</nNF><serie>1</serie><dhEmi>2026-10-01T10:00:00-03:00</dhEmi></ide><emit><CNPJ>12345678000190</CNPJ></emit><dest><xNome>=1+1</xNome></dest><det nItem="1"><prod><cProd>00007</cProd><xProd>Produto de teste</xProd><NCM>01234567</NCM><CFOP>5102</CFOP><vUnCom>0.1234567890</vUnCom><qCom>3.0000</qCom><vProd>0.37</vProd></prod><imposto><ICMS><ICMSSN102><orig>0</orig><CSOSN>102</CSOSN></ICMSSN102></ICMS><PIS><PISQtde><CST>03</CST><qBCProd>3</qBCProd><vAliqProd>0.1234</vAliqProd><vPIS>0.37</vPIS></PISQtde></PIS></imposto></det><total><retTrib><vRetPIS>{ret}</vRetPIS><vRetCOFINS>0.00</vRetCOFINS></retTrib></total></infNFe></NFe></nfeProc>'''.encode()

class ReaderTests(unittest.TestCase):
    def setUp(self):app.reset()
    def test_zero_retention_and_precision(self):
        app.import_xml(xml(),'a.xml');r=app.STATE['items'][0]
        self.assertEqual(len(app.STATE['retained']),0)
        self.assertEqual(r['Vl. Unitario'],'0,123456789')
        self.assertEqual(r['Qtd'],'3')
    def test_quantity_and_csosn(self):
        app.import_xml(xml(),'a.xml');r=app.STATE['items'][0]
        self.assertEqual(r['CSOSN'],'102');self.assertEqual(r['CST ICMS'],'')
        self.assertEqual(r['Base PIS'],'');self.assertEqual(r['Qtd Base PIS'],'3')
        self.assertEqual(r['Aliquota PIS por unidade'],'0,1234')
    def test_positive_retention_and_duplicate(self):
        app.import_xml(xml(ret='1.50'),'a.xml');app.import_xml(xml(ret='1.50'),'b.xml')
        self.assertEqual(app.STATE['notes'],1);self.assertEqual(app.STATE['duplicates'],1)
        self.assertEqual(app.STATE['retained'][0]['Vl. PIS Retido'],'1,50')
    def test_without_namespace(self):
        app.import_xml(xml(namespace=False),'a.xml');self.assertEqual(app.STATE['notes'],1)
    def test_bare_nfe(self):
        from parser import ET
        root=ET.fromstring(xml());nfe=list(root)[0]
        app.import_xml(ET.tostring(nfe),'a.xml');self.assertEqual(app.STATE['notes'],1)
    def test_xml_safety(self):
        for b in [b'<!DOCTYPE n [<!ENTITY a "x">]><n/>','<!DOCTYPE n><n/>'.encode('utf-16')]:
            with self.assertRaisesRegex(ValueError,'DTD'):app.import_xml(b,'a.xml')
        with self.assertRaises(ValueError):app.import_xml(b'<resNFe/>','a.xml')
    def test_csv_and_excel_safety(self):
        app.import_xml(xml(ret='2.00'),'a.xml')
        csv=app.csv_bytes(app.STATE['items'],app.FIELDS_ITENS).decode('utf-8-sig')
        self.assertIn("'=1+1",csv)
        wb=load_workbook(BytesIO(app.xlsx_bytes()))
        ws=wb['Itens'];heads=[c.value for c in ws[1]]
        self.assertEqual(ws.cell(2,heads.index('Destinatario Razao')+1).data_type,'s')
        self.assertEqual(ws.cell(2,heads.index('Codigo')+1).value,'00007')
        self.assertAlmostEqual(ws.cell(2,heads.index('Vl. Unitario')+1).value,.123456789)
        self.assertEqual(ws.freeze_panes,'A2');self.assertEqual(wb.sheetnames,['Itens','Retencoes','Processamento'])
    def test_filter_and_reset(self):
        app.import_xml(xml(),'a.xml');self.assertEqual(len(app.filtered('items','produto')),1)
        self.assertEqual(app.filtered('items','inexistente'),[])
        app.reset();self.assertEqual(app.STATE['items'],[])

class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        cls.url=f'http://127.0.0.1:{cls.server.server_port}'
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
    @classmethod
    def tearDownClass(cls):cls.server.shutdown();cls.server.server_close()
    def request(self,path,body=None,token=True,headers=None):
        h={'X-App-Token':app.TOKEN} if token else {}
        h.update(headers or {})
        return urlopen(Request(self.url+path,data=body,headers=h))
    def test_missing_token(self):
        with self.assertRaises(HTTPError) as ctx:self.request('/api/status',token=False)
        self.assertEqual(ctx.exception.code,403)
    def test_wrong_origin_and_host(self):
        for h in [{'Origin':'https://evil.example'},{'Host':'evil.example'}]:
            with self.assertRaises(HTTPError) as ctx:self.request('/api/reset',b'',headers=h)
            self.assertEqual(ctx.exception.code,403)
    def test_workflow_and_export(self):
        self.request('/api/reset',b'')
        r=json.load(self.request('/api/upload?name=test.xml',xml()))
        self.assertEqual(r['summary']['notes'],1)
        r=json.load(self.request('/api/upload?name=duplicate.xml',xml()))
        self.assertEqual(r['summary']['duplicates'],1)
        r=json.load(self.request('/api/upload?name=invalid.xml',b'<oops>'))
        self.assertEqual(r['summary']['errors'],1)
        r=json.load(self.request('/api/results?kind=items&q=produto'))
        self.assertEqual(r['total'],1)
        data=self.request('/api/export?format=xlsx').read()
        self.assertEqual(load_workbook(BytesIO(data))['Itens'].max_row,2)
        r=self.request('/');self.assertIn('frame-ancestors',r.headers['Content-Security-Policy'])
        self.assertIn(b'Leitor XML',r.read())
if __name__=='__main__':unittest.main()

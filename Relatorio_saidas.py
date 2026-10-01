"""Leitor XML 2.0 — aplicação local, sem envio de documentos a serviços externos."""
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlsplit, parse_qs
from pathlib import Path
from io import BytesIO, StringIO
from threading import RLock, Timer
from decimal import Decimal
import json, csv, secrets, webbrowser, argparse, re
from parser import parse_nfe_xml, FIELDS_ITENS, FIELDS_RETIDO

BASE = Path(__file__).resolve().parent
TOKEN = secrets.token_urlsafe(32)
LOCK = RLock()
MAX_XML = 20 * 1024 * 1024
MAX_SESSION = 250 * 1024 * 1024
STATE = {}

def reset():
    STATE.clear()
    STATE.update(items=[], retained=[], logs=[], keys=set(), notes=0, duplicates=0, errors=0, files=0, bytes=0)
reset()

def summary():
    return {k:STATE[k] for k in ['notes','duplicates','errors','files','bytes']} | {'items':len(STATE['items']), 'retained':len(STATE['retained']), 'xlsx':xlsx_available()}

def xlsx_available():
    try: import openpyxl; return True
    except ImportError: return False

def import_xml(data, name):
    if len(data) > MAX_XML: raise ValueError('XML maior que 20 MB.')
    # Reject DTD/entities even in UTF-16/32 encoded XML.
    probe = data.replace(b'\x00', b'').upper()
    if b'<!DOCTYPE' in probe or b'<!ENTITY' in probe:
        raise ValueError('XML com DTD ou entidades não é permitido.')
    items, retained = parse_nfe_xml(BytesIO(data))
    if items is None: raise ValueError('Este arquivo não contém uma NF-e completa (infNFe).')
    if not items: raise ValueError('NF-e sem itens.')
    key = re.sub(r'\D', '', items[0]['Chave Acesso'])
    if len(key) != 44: raise ValueError('Chave de acesso ausente ou inválida.')
    if key in STATE['keys']:
        STATE['duplicates'] += 1
        return {'file':name,'status':'Duplicada','message':'Chave já importada; arquivo ignorado.'}
    STATE['keys'].add(key)
    STATE['items'].extend(items)
    if retained: STATE['retained'].append(retained)
    STATE['notes'] += 1
    return {'file':name,'status':'Importada','message':f"NF {items[0]['Numero Nota']} · {len(items)} item(ns)"}

def filtered(kind, query=''):
    rows = STATE['retained'] if kind == 'retained' else STATE['logs'] if kind == 'logs' else STATE['items']
    term = query.casefold().strip()
    return [r for r in rows if not term or term in ' '.join(str(v) for v in r.values()).casefold()]

def fields(kind):
    return FIELDS_RETIDO if kind=='retained' else ['file','status','message'] if kind=='logs' else FIELDS_ITENS

def csv_bytes(rows, columns):
    out=StringIO(newline='')
    writer=csv.DictWriter(out,fieldnames=columns,delimiter=';',quoting=csv.QUOTE_ALL)
    writer.writeheader()
    for row in rows:
        safe={k:("'"+str(row.get(k,'')) if str(row.get(k,'')).lstrip().startswith(('=','+','-','@')) else row.get(k,'')) for k in columns}
        writer.writerow(safe)
    return out.getvalue().encode('utf-8-sig')

def xlsx_bytes(query=''):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook(); wb.remove(wb.active)
    for name,kind in [('Itens','items'),('Retencoes','retained'),('Processamento','logs')]:
        ws=wb.create_sheet(name); columns=fields(kind); ws.append(columns)
        for row in filtered(kind,query if kind!='logs' else ''):
            values=[]
            for k in columns:
                value=str(row.get(k,''))
                numeric=k.startswith(('Vl.','Base ','% ','Qtd','Aliquota '))
                if numeric and value:
                    try: value=Decimal(value.replace(',','.'))
                    except Exception: pass
                values.append(value)
            ws.append(values)
            for c in ws[ws.max_row]:
                if isinstance(c.value,str): c.data_type='s'; c.number_format='@'
                elif c.value is not None: c.number_format='#,##0.##########'
        for cell in ws[1]:
            cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='193D65')
            cell.alignment=Alignment(vertical='center')
        ws.row_dimensions[1].height=28
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for i,k in enumerate(columns,1): ws.column_dimensions[get_column_letter(i)].width=48 if k=='Chave Acesso' else 36 if 'Razao' in k or k=='Descricao' or k=='message' else max(14,min(26,len(k)+3))
    out=BytesIO();wb.save(out);return out.getvalue()

class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args): pass
    def respond(self,body,status=200,content_type='application/json; charset=utf-8',filename=None):
        if isinstance(body,dict): body=json.dumps(body,ensure_ascii=False).encode()
        elif isinstance(body,str): body=body.encode()
        self.send_response(status)
        self.send_header('Content-Type',content_type)
        self.send_header('Content-Length',str(len(body)))
        self.send_header('Cache-Control','no-store')
        self.send_header('X-Content-Type-Options','nosniff')
        self.send_header('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; connect-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
        if filename:self.send_header('Content-Disposition',f'attachment; filename="{filename}"')
        self.end_headers();self.wfile.write(body)
    def valid_host(self):
        return self.headers.get('Host')==f'127.0.0.1:{self.server.server_port}'
    def authorized(self):
        return self.valid_host() and secrets.compare_digest(self.headers.get('X-App-Token',''),TOKEN)
    def do_GET(self):
        if not self.valid_host(): return self.respond({'error':'Host não permitido.'},403)
        path=urlsplit(self.path).path
        assets={'/':('index.html','text/html; charset=utf-8'),'/app.js':('app.js','application/javascript'),'/style.css':('style.css','text/css')}
        if path in assets:
            f,ctype=assets[path]; data=(BASE/'web'/f).read_bytes()
            if path=='/':data=data.replace(b'__TOKEN__',TOKEN.encode())
            return self.respond(data,content_type=ctype)
        if path=='/favicon.ico': return self.respond(b'',204,'image/x-icon')
        if not self.authorized():return self.respond({'error':'Sessão inválida. Reabra o aplicativo.'},403)
        q=parse_qs(urlsplit(self.path).query); kind=q.get('kind',['items'])[0];term=q.get('q',[''])[0]
        with LOCK:
            if path=='/api/status':return self.respond(summary())
            if path=='/api/results':
                rows=filtered(kind,term);offset=max(0,int(q.get('offset',['0'])[0]));limit=min(200,max(1,int(q.get('limit',['100'])[0])))
                return self.respond({'rows':rows[offset:offset+limit],'total':len(rows),'fields':fields(kind)})
            if path=='/api/export':
                if q.get('format',['csv'])[0]=='xlsx':
                    if not xlsx_available():return self.respond({'error':'Instale openpyxl para exportar Excel. O CSV continua disponível.'},409)
                    return self.respond(xlsx_bytes(term),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename='Leitor_XML.xlsx')
                return self.respond(csv_bytes(filtered(kind,term),fields(kind)),content_type='text/csv; charset=utf-8',filename=f'relatorio_{kind}.csv')
        return self.respond({'error':'Recurso não encontrado.'},404)
    def do_POST(self):
        if not self.authorized():return self.respond({'error':'Sessão inválida.'},403)
        origin=self.headers.get('Origin')
        if origin and origin!=f'http://127.0.0.1:{self.server.server_port}':return self.respond({'error':'Origem não permitida.'},403)
        path=urlsplit(self.path).path
        try:
            size=int(self.headers.get('Content-Length','0'))
            if size<0 or size>MAX_XML: return self.respond({'error':'Limite de 20 MB por XML.'},413)
            body=self.rfile.read(size)
            with LOCK:
                if path=='/api/reset':reset();return self.respond(summary())
                if path=='/api/upload':
                    if STATE['bytes']+size>MAX_SESSION:return self.respond({'error':'Limite de 250 MB por sessão. Exporte os resultados e inicie outra importação.'},413)
                    name=parse_qs(urlsplit(self.path).query).get('name',['arquivo.xml'])[0][-500:]
                    STATE['files']+=1;STATE['bytes']+=size
                    try: result=import_xml(body,name)
                    except Exception as exc:
                        STATE['errors']+=1;result={'file':name,'status':'Erro','message':str(exc)[:300]}
                    STATE['logs'].append(result)
                    return self.respond({'result':result,'summary':summary()})
            return self.respond({'error':'Recurso não encontrado.'},404)
        except Exception as exc:return self.respond({'error':str(exc)[:200]},400)

def main():
    p=argparse.ArgumentParser(description='Leitor XML local');p.add_argument('--port',type=int,default=8765);p.add_argument('--no-browser',action='store_true');args=p.parse_args()
    try:server=ThreadingHTTPServer(('127.0.0.1',args.port),Handler)
    except OSError:
        server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    url=f'http://127.0.0.1:{server.server_port}'
    print(f'Leitor XML 2.0 | {url}\nMantenha esta janela aberta. Ctrl+C encerra o aplicativo.\nXMLs e resultados ficam em memória; exporte antes de encerrar.')
    if not args.no_browser:Timer(.7,lambda:webbrowser.open(url)).start()
    try:server.serve_forever()
    except KeyboardInterrupt:pass
    finally:server.server_close()
if __name__=='__main__':main()
